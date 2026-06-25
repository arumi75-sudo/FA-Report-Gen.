"""Word(.docx) 문서 생성기 - 현장별 이슈 정리 형태."""
import re
import os
from datetime import datetime
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def _set_cell_bg(cell, hex_color: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _add_heading(doc: Document, text: str, level: int):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    if level == 1:
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x8B, 0x00, 0x00)  # LG 딥레드
    elif level == 2:
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)  # 진청색
    else:
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x40, 0x40, 0x40)
    return p


def _add_bullet(doc: Document, text: str):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    run.font.size = Pt(10)
    return p


def _parse_markdown_section(text: str) -> dict[str, list[str]]:
    """마크다운 보고서를 섹션별로 파싱."""
    sections: dict[str, list[str]] = {}
    current = "기타"
    for line in text.split("\n"):
        line = line.strip()
        if not line:
            continue
        if line.startswith("## ") or line.startswith("### "):
            current = line.lstrip("#").strip()
            sections.setdefault(current, [])
        elif line.startswith("# "):
            continue  # 문서 제목은 skip
        else:
            clean = re.sub(r"^\*{1,2}|^\-\s|^\d+\.\s", "", line).strip("*")
            if clean:
                sections.setdefault(current, []).append(clean)
    return sections


def generate_word_document(content: str, title: str, output_path: str) -> str:
    """마크다운 보고서 내용을 Word 파일로 생성."""
    os.makedirs(os.path.dirname(output_path) if os.path.dirname(output_path) else ".", exist_ok=True)

    doc = Document()

    # 기본 폰트 설정
    style = doc.styles["Normal"]
    style.font.name = "맑은 고딕"
    style.font.size = Pt(10)

    # 헤더: 보고서 제목
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_para.add_run(title)
    title_run.font.size = Pt(18)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(0x8B, 0x00, 0x00)

    # 작성일
    date_para = doc.add_paragraph()
    date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_run = date_para.add_run(f"작성일: {datetime.now().strftime('%Y년 %m월 %d일')}")
    date_run.font.size = Pt(10)
    date_run.font.color.rgb = RGBColor(0x60, 0x60, 0x60)

    doc.add_paragraph()  # 공백

    # 구분선
    hr = doc.add_paragraph("─" * 55)
    hr.paragraph_format.space_after = Pt(8)
    hr.runs[0].font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)

    # 본문 파싱 및 작성
    sections = _parse_markdown_section(content)

    # 현장 구분 처리: "■ ESMI 현장" 같은 패턴이 있으면 섹션별로 분리
    site_pattern = re.compile(r"^[■▶◆\*#]?\s*(ESMI|KSMI|CNMI|PLMI|[A-Z]{2,6}MI)\s*(현장)?", re.IGNORECASE)

    current_site = None
    site_sections: dict[str, list[tuple[str, list[str]]]] = {}  # site → [(section_title, lines)]

    for sec_title, lines in sections.items():
        m = site_pattern.match(sec_title)
        if m:
            current_site = m.group(1).upper()
            site_sections.setdefault(current_site, [])
        elif current_site:
            site_sections[current_site].append((sec_title, lines))
        else:
            # 현장 구분 없는 경우 바로 출력
            _add_heading(doc, sec_title, 2)
            for line in lines:
                _add_bullet(doc, line)

    if site_sections:
        for site_code, sub_sections in site_sections.items():
            _add_heading(doc, f"■ {site_code} 현장", 1)
            for sub_title, lines in sub_sections:
                if lines:
                    _add_heading(doc, f"▶ {sub_title}", 2)
                    for line in lines:
                        _add_bullet(doc, line)
            doc.add_paragraph()
    elif not sections:
        # 구조 없는 텍스트 그대로 넣기
        for line in content.split("\n"):
            line = line.strip()
            if line:
                doc.add_paragraph(line)

    # 푸터
    doc.add_paragraph()
    footer_p = doc.add_paragraph("─" * 55)
    footer_p.runs[0].font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)
    footer_note = doc.add_paragraph("※ 본 문서는 FA Report Generator (Claude AI)에 의해 자동 생성되었습니다.")
    footer_note.runs[0].font.size = Pt(8)
    footer_note.runs[0].font.color.rgb = RGBColor(0x99, 0x99, 0x99)

    doc.save(output_path)
    return output_path


if __name__ == "__main__":
    import sys
    if len(sys.argv) < 3:
        print("사용법: python generate_word.py <input_text_file> <output_docx>")
        sys.exit(1)
    with open(sys.argv[1], encoding="utf-8") as f:
        text = f.read()
    out = generate_word_document(text, "주간 업무 보고", sys.argv[2])
    print(f"생성됨: {out}")
