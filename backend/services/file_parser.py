import io
from pathlib import Path


async def parse_file(filename: str, content: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        return _parse_pdf(content)
    elif ext in (".docx", ".doc"):
        return _parse_docx(content)
    elif ext in (".xlsx", ".xls"):
        return _parse_excel(content)
    elif ext in (".txt", ".md", ".csv"):
        return content.decode("utf-8", errors="replace")
    else:
        raise ValueError(f"지원하지 않는 파일 형식: {ext}")


def _parse_pdf(content: bytes) -> str:
    import fitz  # PyMuPDF
    doc = fitz.open(stream=content, filetype="pdf")
    texts = []
    for page in doc:
        texts.append(page.get_text())
    doc.close()
    return "\n".join(texts)


def _parse_docx(content: bytes) -> str:
    from docx import Document
    doc = Document(io.BytesIO(content))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells)
            if row_text.strip():
                paragraphs.append(row_text)
    return "\n".join(paragraphs)


def _parse_excel(content: bytes) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    texts = []
    for sheet in wb.worksheets:
        texts.append(f"[시트: {sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            row_text = " | ".join(str(cell) if cell is not None else "" for cell in row)
            if row_text.strip(" |"):
                texts.append(row_text)
    wb.close()
    return "\n".join(texts)
